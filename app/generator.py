import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from num2words import num2words
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, PageBreak
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY, TA_RIGHT

OUTPUT_DIR = os.environ.get("OUTPUT_DIR", "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)


def fmt_money(value: float) -> str:
    entero = int(round(value))
    centavos = round((value - int(value)) * 100)
    entero_fmt = f"{entero:,}".replace(",", ".")
    return f"{entero_fmt},{centavos:02d}"


def numero_a_letras(value: float) -> str:
    entero = int(round(value))
    centavos = round((value - int(value)) * 100)
    letras = num2words(entero, lang="es")
    return f"{letras} con {centavos:02d}/100"


def safe_filename(name: str) -> str:
    name = name.replace(" ", "_")
    return re.sub(r"[^a-zA-Z0-9_\-]", "", name)


def _tarea_realizada_items(data: dict):
    tipo = (data.get("tipo_contribuyente") or "monotributista").strip().lower()
    periodo = f"{data['mes_inicio_label']} {data['anio_inicio']} a {data['mes_fin_label']} {data['anio_fin']}"

    if tipo == "responsable_inscripto":
        return [
            f"• Declaraciones Juradas de IVA de los períodos {periodo}.",
            f"• Declaraciones de Actividades Económicas de Salta de los períodos {periodo}.",
        ]

    return ["• Facturas C emitidas en su condición de Monotributista."]


def _manifestacion_table_title(data: dict) -> str:
    tipo = (data.get("tipo_contribuyente") or "monotributista").strip().lower()
    if tipo == "responsable_inscripto":
        return "Ventas netas declaradas en IVA y en Actividades Económicas:"
    return "Facturas C de Ventas (Importes Netos):"


def _tratamiento(data: dict) -> str:
    genero = (data.get("genero") or "masculino").strip().lower()
    return "Señora" if genero == "femenino" else "Señor"


def _set_margins(doc, top=2, bottom=2, left=2.5, right=2.5):
    for section in doc.sections:
        section.top_margin = Cm(top)
        section.bottom_margin = Cm(bottom)
        section.left_margin = Cm(left)
        section.right_margin = Cm(right)


def _add_run(para, text, bold=False, italic=False, underline=False, size=10):
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return run


def _para(doc, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _remove_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for border_name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "none")
        tcBorders.append(border)
    tcPr.append(tcBorders)


# ───────────────────────────────────────────────────────────────
#  WORD 1 – Manifestación de Ingresos
# ───────────────────────────────────────────────────────────────
def generate_word_manifestacion(data: dict, filepath: str):
    doc = Document()
    _set_margins(doc)

    # Título
    p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    _add_run(p, "MANIFESTACION DE INGRESOS", bold=True, underline=True, size=12)

    # Intro cliente
    p = _para(doc, space_after=8)
    _add_run(p, (
        f"{data['nombre_completo']} con DNI N.º {data['dni']} C.U.I.T. N.º {data['cuit']}, "
        f"con domicilio real en {data['domicilio']} de la ciudad de {data['ciudad']} "
        f"provincia de {data['provincia']}."
    ))

    # Declaración
    p = _para(doc, space_after=10)
    _add_run(p, (
        f"DECLARO: que he percibido ingresos Netos correspondientes al período comprendido entre "
        f"los meses de {data['mes_inicio_label']} {data['anio_inicio']} hasta "
        f"{data['mes_fin_label']} {data['anio_fin']}, que ascienden a la suma de pesos "
        f"{numero_a_letras(data['total'])} ($ {fmt_money(data['total'])}), lo que hace un promedio "
        f"mensual de pesos {numero_a_letras(data['promedio'])} ($ {fmt_money(data['promedio'])}). "
        f"Dichos ingresos se conforman de la siguiente manera:"
    ))

    # Subtítulo
    p = _para(doc, space_after=6)
    _add_run(p, _manifestacion_table_title(data), italic=True, underline=True)

    # Tabla montos sin bordes
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Cm(12)
    table.columns[1].width = Cm(4)
    for label, monto in zip(data["month_labels"], data["montos"]):
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = f"$ {fmt_money(monto)}"
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for cell in row.cells:
            _remove_cell_borders(cell)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0, 0, 0)

    # Total
    p = _para(doc, WD_ALIGN_PARAGRAPH.RIGHT, space_before=8, space_after=8)
    _add_run(p, f"TOTAL........................................................$ {fmt_money(data['total'])}", bold=True, size=11)

    # Cierre
    p = _para(doc, space_after=30)
    _add_run(p, (
        f"A los efectos de ser presentada ante quien corresponda, se suscribe la presente en la "
        f"ciudad de {data['ciudad']}, provincia de {data['provincia']}, a los {data['dia_emision']} "
        f"días del mes de {data['mes_emision_label']} de {data['anio_emision_letras']}."
    ))

    p = _para(doc, space_after=0)
    _add_run(p, (
        f"Firmada a efectos de su identificación con mi certificación de fecha "
        f"{data['dia_emision']:02d}/{data['mes_emision']:02d}/{data['anio_emision']}.-"
    ))

    doc.save(filepath)


# ───────────────────────────────────────────────────────────────
#  WORD 2 – Certificación Contable
# ───────────────────────────────────────────────────────────────
def generate_word_certificacion(data: dict, filepath: str):
    doc = Document()
    _set_margins(doc)

    # Título
    p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _add_run(p, "CERTIFICACION CONTABLE SOBRE", bold=True, underline=True, size=12)
    p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    _add_run(p, "MANIFESTACION DE INGRESOS PERSONALES", bold=True, underline=True, size=12)

    # Datos cliente
    for line in [
        f"{_tratamiento(data)} {data['nombre_completo']}",
        f"Domicilio Real: {data['domicilio']} – {data['ciudad']} – {data['provincia']}",
        f"C.U.I.T./CUIL N.º: {data['cuit']}",
        f"Actividad: {data['actividad']}",
    ]:
        p = _para(doc, WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
        _add_run(p, line)
    doc.add_paragraph()

    def section(title, body_lines):
        p = _para(doc, space_after=6)
        _add_run(p, title, bold=True, italic=True)
        for line in body_lines:
            p = _para(doc, space_after=8)
            _add_run(p, line)

    section("Identificación de la información objeto de la certificación", [
        (
            f"He sido contratado para emitir una certificación sobre la manifestación de ingresos preparada "
            f"por {data['nombre_completo']} (en adelante \"el comitente\"), DNI {data['dni']}, dedicado a la "
            f"actividad de {data['actividad']}, referida a sus ingresos por el período transcurrido entre el "
            f"01 de {data['mes_inicio_label'].capitalize()} de {data['anio_inicio']}, y el 31 de "
            f"{data['mes_fin_label'].capitalize()} de {data['anio_fin']}, la cual se adjunta firmada por mí "
            f"al solo efecto de su identificación, para su presentación ante quien corresponda."
        ),
        (
            "La certificación se aplica a ciertas situaciones de hecho o comprobaciones especiales, a través "
            "de la constatación con registros contables y otra documentación de respaldo. Este trabajo "
            "profesional no constituye una auditoría ni una revisión y, por lo tanto, las manifestaciones del "
            "contador público no representan la emisión de un juicio técnico respecto de la información "
            "objeto de la certificación."
        ),
    ])

    section("Responsabilidades del comitente", [
        "El comitente es responsable de la preparación y emisión de la manifestación adjunta que "
        "presenta la información mencionada en la sección precedente."
    ])

    section("Responsabilidades del contador público", [
        "Mi responsabilidad consiste en emitir una certificación sobre la información que se menciona "
        "en la primera sección. He llevado a cabo mi encargo de conformidad con las normas incluidas "
        "en el capítulo VI de la Resolución Técnica N° 37 de la Federación Argentina de Consejos "
        "Profesionales de Ciencias Económicas (FACPCE). Dichas normas exigen que cumpla los "
        "requerimientos de ética, así como que planifique mi tarea. Soy independiente del comitente "
        "y he cumplido las demás responsabilidades de ética de conformidad con los requerimientos del "
        "Código de Ética del Consejo Profesional de Ciencias Económicas y de la Resolución "
        "Técnica N° 37 de la FACPCE."
    ])

    section("Tarea realizada", [
        "Mi tarea profesional se limitó a cotejar la información incluida en la manifestación de ingresos "
        "mencionada en la primera sección de esta certificación con la siguiente documentación:",
    ] + _tarea_realizada_items(data))

    section("Manifestación profesional", [
        "Sobre la base de las tareas descriptas, certifico que la información detallada en la "
        "manifestación de ingresos mencionada en la sección \"Identificación de la información objeto "
        "de la certificación\" concuerda con la documentación respaldatoria señalada en la sección "
        "precedente."
    ])

    p = _para(doc, space_after=0)
    _add_run(p, (
        f"{data['provincia']}, a los {data['dia_emision']} días del mes de "
        f"{data['mes_emision_label']} de {data['anio_emision_letras']}."
    ))

    doc.save(filepath)


# ───────────────────────────────────────────────────────────────
#  PDF – Manifestación + Certificación (2 páginas)
# ───────────────────────────────────────────────────────────────
def generate_pdf(data: dict, filepath: str):
    doc = SimpleDocTemplate(
        filepath, pagesize=A4,
        rightMargin=2.5 * cm, leftMargin=2.5 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm,
    )

    normal = ParagraphStyle("n", fontName="Helvetica", fontSize=10,
                             leading=14, alignment=TA_JUSTIFY, spaceAfter=8)
    title_s = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=12,
                              alignment=TA_CENTER, spaceAfter=4, leading=16)
    bi = ParagraphStyle("bi", fontName="Helvetica-BoldOblique", fontSize=10,
                         leading=14, alignment=TA_JUSTIFY, spaceAfter=6)
    right_s = ParagraphStyle("r", fontName="Helvetica-Bold", fontSize=11,
                              alignment=TA_RIGHT, spaceAfter=6)
    left_s = ParagraphStyle("l", fontName="Helvetica", fontSize=10,
                              leading=14, alignment=TA_LEFT, spaceAfter=4)

    story = []

    # ── Página 1: Certificación ──
    story.append(Paragraph("<u><b>CERTIFICACION CONTABLE SOBRE</b></u>", title_s))
    story.append(Paragraph("<u><b>MANIFESTACION DE INGRESOS PERSONALES</b></u>", title_s))
    story.append(Spacer(1, 0.3 * cm))
    for line in [
        f"{_tratamiento(data)} {data['nombre_completo']}",
        f"Domicilio Real: {data['domicilio']} – {data['ciudad']} – {data['provincia']}",
        f"C.U.I.T./CUIL N.º: {data['cuit']}",
        f"Actividad: {data['actividad']}",
    ]:
        story.append(Paragraph(line, left_s))
    story.append(Spacer(1, 0.3 * cm))

    def pdf_section(titulo, cuerpo):
        story.append(Paragraph(f"<b><i>{titulo}</i></b>", bi))
        for p in cuerpo:
            story.append(Paragraph(p, normal))

    pdf_section("Identificación de la información objeto de la certificación", [
        (f"He sido contratado para emitir una certificación sobre la manifestación de ingresos preparada "
         f"por {data['nombre_completo']} (en adelante \"el comitente\"), DNI {data['dni']}, dedicado a la "
         f"actividad de {data['actividad']}, referida a sus ingresos por el período transcurrido entre el "
         f"01 de {data['mes_inicio_label'].capitalize()} de {data['anio_inicio']}, y el 31 de "
         f"{data['mes_fin_label'].capitalize()} de {data['anio_fin']}, la cual se adjunta firmada por mí "
         f"al solo efecto de su identificación, para su presentación ante quien corresponda."),
        ("La certificación se aplica a ciertas situaciones de hecho o comprobaciones especiales, a través "
         "de la constatación con registros contables y otra documentación de respaldo. Este trabajo "
         "profesional no constituye una auditoría ni una revisión y, por lo tanto, las manifestaciones del "
         "contador público no representan la emisión de un juicio técnico respecto de la información "
         "objeto de la certificación."),
    ])
    pdf_section("Responsabilidades del comitente", [
        "El comitente es responsable de la preparación y emisión de la manifestación adjunta que "
        "presenta la información mencionada en la sección precedente."
    ])
    pdf_section("Responsabilidades del contador público", [
        "Mi responsabilidad consiste en emitir una certificación sobre la información que se menciona "
        "en la primera sección. He llevado a cabo mi encargo de conformidad con las normas incluidas "
        "en el capítulo VI de la Resolución Técnica N° 37 de la Federación Argentina de Consejos "
        "Profesionales de Ciencias Económicas (FACPCE). Dichas normas exigen que cumpla los "
        "requerimientos de ética, así como que planifique mi tarea. Soy independiente del comitente "
        "y he cumplido las demás responsabilidades de ética de conformidad con los requerimientos del "
        "Código de Ética del Consejo Profesional de Ciencias Económicas y de la Resolución "
        "Técnica N° 37 de la FACPCE."
    ])
    pdf_section("Tarea realizada", [
        "Mi tarea profesional se limitó a cotejar la información incluida en la manifestación de ingresos "
        "mencionada en la primera sección de esta certificación con la siguiente documentación:",
    ] + _tarea_realizada_items(data))
    pdf_section("Manifestación profesional", [
        "Sobre la base de las tareas descriptas, certifico que la información detallada en la "
        "manifestación de ingresos mencionada en la sección \"Identificación de la información objeto "
        "de la certificación\" concuerda con la documentación respaldatoria señalada en la sección "
        "precedente."
    ])
    story.append(Spacer(1, 0.3 * cm))
    story.append(Paragraph(
        f"{data['provincia']}, a los {data['dia_emision']} días del mes de "
        f"{data['mes_emision_label']} de {data['anio_emision_letras']}.", normal))

    # ── Página 2: Manifestación ──
    story.append(PageBreak())
    story.append(Paragraph("<u><b>MANIFESTACION DE INGRESOS</b></u>", title_s))
    story.append(Spacer(1, 0.3 * cm))
    story.append(Paragraph(
        f"{data['nombre_completo']} con DNI N.º {data['dni']} C.U.I.T. N.º {data['cuit']}, "
        f"con domicilio real en {data['domicilio']} de la ciudad de {data['ciudad']} "
        f"provincia de {data['provincia']}.", normal))
    story.append(Paragraph(
        f"DECLARO: que he percibido ingresos Netos correspondientes al período comprendido entre "
        f"los meses de {data['mes_inicio_label']} {data['anio_inicio']} hasta "
        f"{data['mes_fin_label']} {data['anio_fin']}, que ascienden a la suma de pesos "
        f"{numero_a_letras(data['total'])} ($ {fmt_money(data['total'])}), lo que hace un promedio "
        f"mensual de pesos {numero_a_letras(data['promedio'])} ($ {fmt_money(data['promedio'])}). "
        f"Dichos ingresos se conforman de la siguiente manera:", normal))
    story.append(Paragraph(f"<i><u>{_manifestacion_table_title(data)}</u></i>", normal))
    story.append(Spacer(1, 0.2 * cm))

    table_data = [[label, f"$ {fmt_money(monto)}"]
                  for label, monto in zip(data["month_labels"], data["montos"])]
    t = Table(table_data, colWidths=[13 * cm, 4 * cm])
    t.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("LINEBELOW", (0, -1), (-1, -1), 0, colors.white),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.2 * cm))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.black))
    story.append(Paragraph(f"<b>TOTAL $ {fmt_money(data['total'])}</b>", right_s))
    story.append(Spacer(1, 0.5 * cm))
    story.append(Paragraph(
        f"A los efectos de ser presentada ante quien corresponda, se suscribe la presente en la "
        f"ciudad de {data['ciudad']}, provincia de {data['provincia']}, a los {data['dia_emision']} "
        f"días del mes de {data['mes_emision_label']} de {data['anio_emision_letras']}.", normal))
    story.append(Spacer(1, 1.5 * cm))
    story.append(Paragraph(
        f"Firmada a efectos de su identificación con mi certificación de fecha "
        f"{data['dia_emision']:02d}/{data['mes_emision']:02d}/{data['anio_emision']}.-", normal))

    doc.build(story)


# ───────────────────────────────────────────────────────────────
#  WORD 3 – Certificación + Manifestación (1 solo archivo)
# ───────────────────────────────────────────────────────────────
def generate_word_combined(data: dict, filepath: str):
    """Genera un único Word con Certificación (p.1) + Manifestación (p.2)"""
    doc = Document()
    _set_margins(doc)

    # ═══ PÁGINA 1: CERTIFICACIÓN ═══
    p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _add_run(p, "CERTIFICACION CONTABLE SOBRE", bold=True, underline=True, size=12)
    p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    _add_run(p, "MANIFESTACION DE INGRESOS PERSONALES", bold=True, underline=True, size=12)

    # Datos cliente
    for line in [
        f"{_tratamiento(data)} {data['nombre_completo']}",
        f"Domicilio Real: {data['domicilio']} – {data['ciudad']} – {data['provincia']}",
        f"C.U.I.T./CUIL N.º: {data['cuit']}",
        f"Actividad: {data['actividad']}",
    ]:
        p = _para(doc, WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
        _add_run(p, line)
    doc.add_paragraph()

    def section(title, body_lines):
        p = _para(doc, space_after=6)
        _add_run(p, title, bold=True, italic=True)
        for line in body_lines:
            p = _para(doc, space_after=8)
            _add_run(p, line)

    section("Identificación de la información objeto de la certificación", [
        (
            f"He sido contratado para emitir una certificación sobre la manifestación de ingresos preparada "
            f"por {data['nombre_completo']} (en adelante \"el comitente\"), DNI {data['dni']}, dedicado a la "
            f"actividad de {data['actividad']}, referida a sus ingresos por el período transcurrido entre el "
            f"01 de {data['mes_inicio_label'].capitalize()} de {data['anio_inicio']}, y el 31 de "
            f"{data['mes_fin_label'].capitalize()} de {data['anio_fin']}, la cual se adjunta firmada por mí "
            f"al solo efecto de su identificación, para su presentación ante quien corresponda."
        ),
        (
            "La certificación se aplica a ciertas situaciones de hecho o comprobaciones especiales, a través "
            "de la constatación con registros contables y otra documentación de respaldo. Este trabajo "
            "profesional no constituye una auditoría ni una revisión y, por lo tanto, las manifestaciones del "
            "contador público no representan la emisión de un juicio técnico respecto de la información "
            "objeto de la certificación."
        ),
    ])

    section("Responsabilidades del comitente", [
        "El comitente es responsable de la preparación y emisión de la manifestación adjunta que "
        "presenta la información mencionada en la sección precedente."
    ])

    section("Responsabilidades del contador público", [
        "Mi responsabilidad consiste en emitir una certificación sobre la información que se menciona "
        "en la primera sección. He llevado a cabo mi encargo de conformidad con las normas incluidas "
        "en el capítulo VI de la Resolución Técnica N° 37 de la Federación Argentina de Consejos "
        "Profesionales de Ciencias Económicas (FACPCE). Dichas normas exigen que cumpla los "
        "requerimientos de ética, así como que planifique mi tarea. Soy independiente del comitente "
        "y he cumplido las demás responsabilidades de ética de conformidad con los requerimientos del "
        "Código de Ética del Consejo Profesional de Ciencias Económicas y de la Resolución "
        "Técnica N° 37 de la FACPCE."
    ])

    section("Tarea realizada", [
        "Mi tarea profesional se limitó a cotejar la información incluida en la manifestación de ingresos "
        "mencionada en la primera sección de esta certificación con la siguiente documentación:",
    ] + _tarea_realizada_items(data))

    section("Manifestación profesional", [
        "Sobre la base de las tareas descriptas, certifico que la información detallada en la "
        "manifestación de ingresos mencionada en la sección \"Identificación de la información objeto "
        "de la certificación\" concuerda con la documentación respaldatoria señalada en la sección "
        "precedente."
    ])

    p = _para(doc, space_after=0)
    _add_run(p, (
        f"{data['provincia']}, a los {data['dia_emision']} días del mes de "
        f"{data['mes_emision_label']} de {data['anio_emision_letras']}."
    ))

    # ═══ PAGE BREAK ═══
    doc.add_page_break()

    # ═══ PÁGINA 2: MANIFESTACIÓN ═══
    p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    _add_run(p, "MANIFESTACION DE INGRESOS", bold=True, underline=True, size=12)

    # Intro cliente
    p = _para(doc, space_after=8)
    _add_run(p, (
        f"{data['nombre_completo']} con DNI N.º {data['dni']} C.U.I.T. N.º {data['cuit']}, "
        f"con domicilio real en {data['domicilio']} de la ciudad de {data['ciudad']} "
        f"provincia de {data['provincia']}."
    ))

    # Declaración
    p = _para(doc, space_after=10)
    _add_run(p, (
        f"DECLARO: que he percibido ingresos Netos correspondientes al período comprendido entre "
        f"los meses de {data['mes_inicio_label']} {data['anio_inicio']} hasta "
        f"{data['mes_fin_label']} {data['anio_fin']}, que ascienden a la suma de pesos "
        f"{numero_a_letras(data['total'])} ($ {fmt_money(data['total'])}), lo que hace un promedio "
        f"mensual de pesos {numero_a_letras(data['promedio'])} ($ {fmt_money(data['promedio'])}). "
        f"Dichos ingresos se conforman de la siguiente manera:"
    ))

    # Subtítulo
    p = _para(doc, space_after=6)
    _add_run(p, _manifestacion_table_title(data), italic=True, underline=True)

    # Tabla montos sin bordes
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Cm(12)
    table.columns[1].width = Cm(4)
    for label, monto in zip(data["month_labels"], data["montos"]):
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = f"$ {fmt_money(monto)}"
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for cell in row.cells:
            _remove_cell_borders(cell)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0, 0, 0)

    # Total
    p = _para(doc, WD_ALIGN_PARAGRAPH.RIGHT, space_before=8, space_after=8)
    _add_run(p, f"TOTAL........................................................$ {fmt_money(data['total'])}", bold=True, size=11)

    # Cierre
    p = _para(doc, space_after=30)
    _add_run(p, (
        f"A los efectos de ser presentada ante quien corresponda, se suscribe la presente en la "
        f"ciudad de {data['ciudad']}, provincia de {data['provincia']}, a los {data['dia_emision']} "
        f"días del mes de {data['mes_emision_label']} de {data['anio_emision_letras']}."
    ))

    p = _para(doc, space_after=0)
    _add_run(p, (
        f"Firmada a efectos de su identificación con mi certificación de fecha "
        f"{data['dia_emision']:02d}/{data['mes_emision']:02d}/{data['anio_emision']}.-"
    ))

    doc.save(filepath)


# ───────────────────────────────────────────────────────────────
#  Punto de entrada
# ───────────────────────────────────────────────────────────────
def generate_documents(data: dict):
    safe_name = safe_filename(data["nombre_completo"])
    base = f"{safe_name}_{data['anio_fin']}"

    word_path = os.path.join(OUTPUT_DIR, f"{base}_certificacion_manifestacion.docx")
    pdf_path = os.path.join(OUTPUT_DIR, f"{base}_completo.pdf")

    generate_word_combined(data, word_path)
    generate_pdf(data, pdf_path)

    return word_path, pdf_path
